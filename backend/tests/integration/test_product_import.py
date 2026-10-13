"""Bulk product import from CSV / Excel: preview, all-or-nothing import, updates, stock, barcodes, suppliers."""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.category import Category
from app.models.enums import InventoryTxnType
from app.models.inventory import Inventory, InventoryTransaction
from app.models.product import Product
from app.models.supplier import Supplier
from app.services.product_import_service import COLUMNS
from tests.factories import make_product

pytestmark = pytest.mark.integration

HEADER = "name,price,tax_rate,mrp,uom,category,sku,barcodes,is_variable_weight,plu_code,supplier_code,opening_stock,expiry_date"


def _csv(*lines: str) -> bytes:
    return ("\n".join((HEADER,) + lines) + "\n").encode("utf-8")


def _post(client, headers, url, content: bytes, filename="products.csv", **form):
    data = {"file": (io.BytesIO(content), filename), **{k: str(v) for k, v in form.items()}}
    return client.post(url, data=data, headers=headers, content_type="multipart/form-data")


def _preview(client, headers, content, **form):
    return _post(client, headers, "/api/v1/products/import/preview", content, **form)


def _import(client, headers, content, **form):
    return _post(client, headers, "/api/v1/products/import", content, **form)


@pytest.fixture
def supplier(app):
    s = Supplier(supplier_code="SUP-ABC123", company_name="Sharma Traders")
    db.session.add(s)
    db.session.commit()
    return s


def _product(name):
    return db.session.scalar(db.select(Product).where(Product.name == name))


def test_template_download(client, admin_headers):
    resp = client.get("/api/v1/products/import/template", headers=admin_headers)
    assert resp.status_code == 200 and "spreadsheetml" in resp.content_type
    wb = load_workbook(io.BytesIO(resp.data))
    assert [c.value for c in wb["Products"][1]] == COLUMNS
    assert wb["Products"]["I2"].number_format == "@"  # barcodes stay text
    assert "Instructions" in wb.sheetnames


def test_csv_creates_products_with_stock_barcodes_and_categories(client, admin_headers, store, supplier):
    content = _csv(
        "Parle-G 100g,9.50,5,10,pcs,Biscuits,,8901719101038,no,,SUP-ABC123,50,31-12-2026",
        "Tata Salt 1kg,26,5,28,,Grocery,,,,,,20,",
        "साखर (सुटी),42,0,,kg,Loose,,,yes,00123,,100,",
    )
    preview = _preview(client, admin_headers, content, store_id=store.id).get_json()
    assert (preview["creates"], preview["updates"], preview["errors"]) == (3, 0, 0)
    assert preview["new_categories"] == ["Biscuits", "Grocery", "Loose"]
    assert db.session.scalar(db.select(db.func.count(Product.id))) == 0  # preview saves nothing

    result = _import(client, admin_headers, content, store_id=store.id)
    assert result.status_code == 201, result.get_json()
    body = result.get_json()
    assert (body["created"], body["updated"], body["opening_stock_rows"]) == (3, 0, 3)

    parle = _product("Parle-G 100g")
    assert parle.supplier_id == supplier.id and parle.store_id is None
    assert [b.barcode for b in parle.barcodes] == ["8901719101038"] and parle.barcodes[0].is_primary
    assert parle.category.name == "Biscuits" and parle.sku.startswith("SKU-")

    salt = _product("Tata Salt 1kg")
    assert [g["product_id"] for g in body["generated_barcodes"]] == [salt.id]
    assert salt.barcodes[0].barcode.startswith("290")  # in-house code, for stickers

    sugar = _product("साखर (सुटी)")
    assert sugar.is_variable_weight and sugar.plu_code == "123" and sugar.uom.value == "kg"
    assert sugar.barcodes == []  # loose items use scale stickers, no generated barcode

    stock = db.session.scalar(db.select(Inventory).where(Inventory.product_id == parle.id, Inventory.store_id == store.id))
    assert stock.qty_on_hand == Decimal("50")
    txn = db.session.scalar(db.select(InventoryTransaction).where(InventoryTransaction.product_id == parle.id))
    assert txn.txn_type is InventoryTxnType.ADJUSTMENT and txn.notes == "Opening stock (bulk import)"
    assert txn.expiry_date == date(2026, 12, 31)

    # Billing scan finds the imported barcode.
    scan = client.get("/api/v1/products/scan", query_string={"barcode": "8901719101038"}, headers=admin_headers)
    assert scan.get_json()["product"]["id"] == parle.id


def test_xlsx_with_numbers_and_dates(client, admin_headers, store):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Price", "GST", "Barcode", "Stock", "Expiry"])  # aliases, any case
    ws.append(["Maggi 70g", 12, 18, 8901058851120, 30, date(2027, 1, 15)])  # numeric barcode cell keeps every digit
    out = io.BytesIO()
    wb.save(out)
    resp = _import(client, admin_headers, out.getvalue(), filename="list.xlsx", store_id=store.id)
    assert resp.status_code == 201, resp.get_json()
    maggi = _product("Maggi 70g")
    assert maggi.tax_rate == Decimal("18.00") and maggi.barcodes[0].barcode == "8901058851120"
    txn = db.session.scalar(db.select(InventoryTransaction).where(InventoryTransaction.product_id == maggi.id))
    assert txn.qty == Decimal("30") and txn.expiry_date == date(2027, 1, 15)


def test_existing_products_are_updated_not_duplicated(client, admin_headers, store):
    biscuit = make_product(store, name="Biscuit", price="10", tax_rate="5", sku="BIS-1", barcode="1111111111111", mrp="12", stock=7)
    soap = make_product(store, name="Soap", price="30", tax_rate="18", barcode="2222222222222")
    content = _csv(
        ",11,,,,,BIS-1,3333333333333,,,,500,",   # by SKU: new price, MRP kept, extra barcode, stock ignored
        "Soap Big,32,,40,,,,2222222222222,,,,,",  # by barcode: renamed, new price + MRP
    )
    preview = _preview(client, admin_headers, content, store_id=store.id).get_json()
    assert (preview["creates"], preview["updates"], preview["errors"]) == (0, 2, 0)
    assert "Opening stock ignored" in preview["rows"][0]["warnings"][0]
    assert preview["rows"][0]["name"] == "Biscuit" and preview["rows"][0]["product_id"] == biscuit.id

    assert _import(client, admin_headers, content, store_id=store.id).get_json()["updated"] == 2
    db.session.refresh(biscuit)
    db.session.refresh(soap)
    assert biscuit.price == Decimal("11.00") and biscuit.mrp == Decimal("12.00") and biscuit.tax_rate == Decimal("5.00")
    assert sorted(b.barcode for b in biscuit.barcodes) == ["1111111111111", "3333333333333"]
    assert db.session.scalar(db.select(Inventory.qty_on_hand).where(Inventory.product_id == biscuit.id)) == Decimal("7")
    assert (soap.name, soap.price, soap.mrp) == ("Soap Big", Decimal("32.00"), Decimal("40.00"))
    assert db.session.scalar(db.select(db.func.count(Product.id))) == 2


def test_errors_block_the_whole_import(client, admin_headers, store, supplier):
    other = make_product(store, name="Other", sku="OTH-1", barcode="4444444444444")
    make_product(store, name="Loose rice", sku="RICE", is_variable_weight=True)
    content = _csv(
        "Good row,10,5,,,,,5555555555555,,,,,",
        "No price,,5,,,,,,,,,,",
        "Excel broke it,10,5,,,,,8.90172E+12,,,,,",
        "Above MRP,100,18,110,,,,,,,,,",
        "Bad supplier,10,0,,,,,,,,SUP-NOPE,,",
        "Bad unit,10,0,,dozen,,,,,,,,",
        "Dup barcode,10,0,,,,,5555555555555,,,,,",
        "Mixed,10,0,,,,OTH-1,9999999999999|1234567890123,,,,,",
        "Packet with PLU,10,0,,,,,,no,55,,,",
        "Half packet,10,0,,,,,,,,,1.5,",
    )
    # A row claiming another product's barcode under a different SKU.
    content += b"Clash,10,0,,,,NEW-SKU,4444444444444,,,,,\n"
    preview = _preview(client, admin_headers, content, store_id=store.id).get_json()
    errors = {r["name"]: " ".join(r["errors"]) for r in preview["rows"]}
    assert errors["Good row"] == ""
    assert "price is required" in errors["No price"]
    assert "changed by Excel" in errors["Excel broke it"]
    assert "MRP" in errors["Above MRP"]
    assert "SUP-NOPE not found" in errors["Bad supplier"]
    assert "uom" in errors["Bad unit"]
    assert "also on row 2" in errors["Dup barcode"]
    assert "plu_code is only for loose items" in errors["Packet with PLU"]
    assert "whole number" in errors["Half packet"]
    assert "Other" in errors["Clash"]
    assert preview["errors"] == 9

    resp = _import(client, admin_headers, content, store_id=store.id)
    assert resp.status_code == 400 and resp.get_json()["preview"]["errors"] == 9
    assert _product("Good row") is None  # all or nothing
    assert db.session.scalar(db.select(db.func.count(Category.id))) == 0


def test_opening_stock_needs_a_store(client, admin_headers):
    preview = _preview(client, admin_headers, _csv("Pen,5,0,,,,,,,,,10,")).get_json()
    assert "Choose the store" in preview["rows"][0]["errors"][0]
    ok = _preview(client, admin_headers, _csv("Pen,5,0,,,,,,,,,,")).get_json()  # no stock: no store needed
    assert ok["errors"] == 0


def test_store_only_products(client, admin_headers, store):
    resp = _import(client, admin_headers, _csv("Local sweets,200,5,,,,,,,,,,"), store_id=store.id, sold_in="store")
    assert resp.status_code == 201
    assert _product("Local sweets").store_id == store.id


def test_file_problems(client, admin_headers):
    assert _preview(client, admin_headers, b"").status_code == 400
    assert _post(client, admin_headers, "/api/v1/products/import/preview", b"x", filename="a.pdf").status_code == 400
    assert _post(client, admin_headers, "/api/v1/products/import/preview", b"x", filename="a.xls").status_code == 400
    assert _preview(client, admin_headers, b"foo,bar\n1,2\n").status_code == 400  # no known headings
    preview = _preview(client, admin_headers, b"name,price,colour\nPen,5,blue\n").get_json()
    assert "colour" in preview["warnings"][0] and preview["creates"] == 1


def test_only_admin_and_manager(client, cashier_headers):
    assert _preview(client, cashier_headers, _csv("Pen,5,0,,,,,,,,,,")).status_code == 403
    assert client.get("/api/v1/products/import/template", headers=cashier_headers).status_code == 403


def test_product_form_supplier_link(client, admin_headers, supplier):
    created = client.post("/api/v1/products", json={"name": "Dal", "price": "90", "supplier_id": supplier.id}, headers=admin_headers)
    assert created.status_code == 201 and created.get_json()["supplier_id"] == supplier.id
    pid = created.get_json()["id"]
    assert client.put(f"/api/v1/products/{pid}", json={"supplier_id": None}, headers=admin_headers).get_json()["supplier_id"] is None
    assert client.post("/api/v1/products", json={"name": "X", "price": "1", "supplier_id": 9999}, headers=admin_headers).status_code == 400
